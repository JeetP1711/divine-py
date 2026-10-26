"""Tests for the in-memory Excel workbook."""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.analytics import (
    calculate_line_metrics,
    daily_summary,
    summarize_by_marketplace,
    summarize_by_sku,
)
from src.exporter import build_excel_report


def test_excel_report_contains_all_sheets_and_calculated_values():
    rows = pd.DataFrame(
        {
            "date": [pd.Timestamp("2026-01-02").date()],
            "order_id": ["=2+3"],
            "marketplace": ["Sample Shop"],
            "sku": ["SAMPLE-1"],
            "product_name": ["Fictional Garden Blend"],
            "units": [2],
            "selling_price": [100],
            "unit_cost": [30],
            "platform_fee": [10],
            "shipping_cost": [5],
            "ad_cost": [2],
            "refund_amount": [0],
        }
    )
    sales = calculate_line_metrics(rows)
    report = build_excel_report(
        sales,
        summarize_by_sku(sales),
        summarize_by_marketplace(sales),
        daily_summary(sales),
        [],
    )
    workbook = load_workbook(BytesIO(report), data_only=True)

    assert workbook.sheetnames == [
        "Sales Details",
        "Product Summary",
        "Marketplace Summary",
        "Daily Summary",
        "Business Alerts",
    ]
    product_sheet = workbook["Product Summary"]
    headers = [cell.value for cell in product_sheet[1]]
    assert "Estimated Contribution Profit" in headers
    assert product_sheet.cell(row=2, column=headers.index("Estimated Contribution Profit") + 1).value == 123
    assert product_sheet.freeze_panes == "A2"
    assert workbook["Business Alerts"]["A2"].value == "No alerts for the selected period."
    sales_sheet = workbook["Sales Details"]
    order_id_column = [cell.value for cell in sales_sheet[1]].index("Order Id") + 1
    assert sales_sheet.cell(row=2, column=order_id_column).value == "=2+3"
    assert sales_sheet.cell(row=2, column=order_id_column).data_type == "s"


def test_excel_report_formats_currency_columns():
    data = pd.DataFrame({"net_revenue": [1234.5], "estimated_profit": [123]})
    report = build_excel_report(data, data, data, data, [])
    workbook = load_workbook(BytesIO(report))
    assert workbook["Sales Details"]["A2"].number_format.startswith("₹")
    assert workbook["Sales Details"]["B2"].number_format.startswith("₹")
