"""In-memory Excel report generation."""

from io import BytesIO
from typing import Dict, List

import pandas as pd
from openpyxl.styles import Font, PatternFill


def build_excel_report(
    sales_details: pd.DataFrame,
    product_summary: pd.DataFrame,
    marketplace_summary: pd.DataFrame,
    daily_summary: pd.DataFrame,
    alerts: List[Dict[str, str]],
) -> bytes:
    """Build the five-sheet report without writing uploaded data to disk."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sales_details.to_excel(writer, sheet_name="Sales Details", index=False)
        product_summary.to_excel(writer, sheet_name="Product Summary", index=False)
        marketplace_summary.to_excel(writer, sheet_name="Marketplace Summary", index=False)
        daily_summary.to_excel(writer, sheet_name="Daily Summary", index=False)
        alerts_frame = pd.DataFrame(alerts)
        if alerts_frame.empty:
            alerts_frame = pd.DataFrame({"Note": ["No alerts for the selected period."]})
        alerts_frame.to_excel(writer, sheet_name="Business Alerts", index=False)

        column_titles = {
            "estimated_profit": "Estimated Contribution Profit",
            "net_revenue": "Net Revenue",
            "profit_margin_pct": "Contribution Margin (%)",
            "contribution_margin_pct": "Contribution Margin (%)",
            "units_sold": "Units Sold",
            "total_refund": "Refund Amount",
            "refund_amount": "Refund Amount",
        }
        currency_columns = {
            "gross_sales",
            "revenue",
            "net_revenue",
            "selling_price",
            "unit_cost",
            "cost_of_goods",
            "platform_fee",
            "shipping_cost",
            "ad_cost",
            "refund_amount",
            "total_refund",
            "estimated_profit",
        }
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.lstrip(" \t\r\n").startswith(
                        ("=", "+", "-", "@")
                    ):
                        cell.data_type = "s"
            for cell in worksheet[1]:
                raw_title = str(cell.value)
                cell.value = column_titles.get(raw_title, raw_title.replace("_", " ").title())
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2E7D32")
            for column_cells in worksheet.columns:
                title = str(column_cells[0].value)
                normalized_title = next(
                    (key for key, label in column_titles.items() if label == title),
                    title.lower().replace(" ", "_"),
                )
                for cell in column_cells[1:]:
                    if normalized_title in currency_columns:
                        cell.number_format = '₹#,##,##0.00;[Red]-₹#,##,##0.00'
                    elif normalized_title.endswith("_pct"):
                        cell.number_format = '0.0"%"'
                    elif normalized_title == "date":
                        cell.number_format = "dd-mmm-yyyy"
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12),
                    38,
                )
    output.seek(0)
    return output.getvalue()
